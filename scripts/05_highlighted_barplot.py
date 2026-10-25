from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
OLEATE_FILE = BASE_DIR / "data" / "plots" / "Oleate_relative_abundance.xlsx"
STEARATE_FILE = BASE_DIR / "data" / "plots" / "Stearate_relative_abundance.xlsx"
OUTPUT_DIR = BASE_DIR / "results" / "05_figures"
OUTPUT_FILE = OUTPUT_DIR / "highlighted_abundance.svg"

LABELS = {
    "K2R3X9": "mer", "K2QYM4": "mcrA", "K2RAN9": "mcrB", "K2R2H1": "mcrG",
    "K2QYN3": "mvhD", "K2RRN3": "mvhA", "K2R6I7": "hdrA2", "K2REI5": "frhG",
    "K2R2M8": "frhB", "K2QFA9": "frhA", "K2RSU0": "cdhA", "K2RBB2": "hmd",
    "K2R351": "mtd", "K2QBS0": "mtrH",
    "A0A0E4GEM6": "etfB", "A0A0E4GCA0": "etfA", "A0A0E4G9L3": "fadA",
    "A0A0E3W2J0": "atpA", "A0A0E4G9E7": "atpD", "A0A0E3W3R4": "ppdK", "A0A0E4C836": "SDR",
}


def read_excel(path):
    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    rows = []
    for row in ws.iter_rows(min_row=2):
        rows.append({
            "accession": row[col["Main Accession"]].value,
            "organism": row[col["Taxonomy"]].value,
            "abundance": row[col["Relative Abundance [%]"]].value,
            "validation": row[col["Validation"]].value,
            "highlighted": row[col["Description"]].fill.fill_type is not None,
        })
    return pd.DataFrame(rows).set_index("accession")


def build_table():
    ole, ste = read_excel(OLEATE_FILE), read_excel(STEARATE_FILE)
    selected = sorted(set(ole.index[ole.highlighted]) | set(ste.index[ste.highlighted]))
    table = pd.DataFrame(index=selected)
    for name, df in (("oleate", ole), ("stearate", ste)):
        table[name] = df["abundance"].reindex(selected)
        table["val_" + name] = df["validation"].reindex(selected)
    organisms = pd.concat([ole["organism"], ste["organism"]]).groupby(level=0).first()
    table["species"] = organisms.reindex(selected).map(
        lambda s: "M. formicicum" if "Methanobacterium" in s else "S. zehnderi")
    table["label"] = [LABELS.get(a, a) for a in selected]
    table["rank"] = table["oleate"].fillna(-1) * 1000 + table["stearate"].fillna(-1)
    return table.sort_values(["species", "rank"], ascending=[True, False])


def plot_highlighted(table, path):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        import seaborn as sns
        from matplotlib.patches import Patch
    except ImportError:
        print("matplotlib or seaborn is not installed: chart not generated")
        return

    plt.rcParams["axes.labelsize"] = 11
    plt.rcParams["xtick.labelsize"] = 10
    plt.rcParams["ytick.labelsize"] = 10
    plt.rcParams["hatch.linewidth"] = 0.5
    sns.set_theme(
        style="whitegrid",
        rc={
            "font.family": ["Linux Libertine", "Libertinus Serif", "DejaVu Serif"],
            "svg.fonttype": "none",
        },
    )
    sns.set_context("paper")

    colors = {"oleate": "#4c78a8", "stearate": "#e45756"}
    fig, ax = plt.subplots(figsize=(6.6, 8.6), dpi=150)
    xmax = table[["oleate", "stearate"]].max().max()

    y, ticks, labels, spans = 0.0, [], [], []
    for species, group in table.groupby("species", sort=False):
        start = y
        for acc, row in group.iterrows():
            for offset, sample in ((-0.2, "oleate"), (0.2, "stearate")):
                value, validation = row[sample], row["val_" + sample]
                yi = y + offset
                if pd.isna(value):
                    ax.text(xmax * 0.012, yi, "n.i.", ha="left", va="center", fontsize=6.5, color="gray", style="italic")
                    continue
                ax.barh(
                    yi, value, height=0.4, color=colors[sample], edgecolor="black", linewidth=0.4,
                    hatch="////" if validation == "Doubtful" else None,
                )
                if value == 0:
                    ax.text(xmax * 0.012, yi, "n.q.", ha="left", va="center", fontsize=6.5, color="gray", style="italic")
                else:
                    ax.text(value + xmax * 0.012, yi, f"{value:.1f}", ha="left", va="center", fontsize=6.5)
            ticks.append(y)
            labels.append(row["label"])
            y += 1.0
        spans.append((start - 0.5, y - 0.5, species))
        y += 0.6

    ax.set_yticks(ticks, labels, style="italic")
    ax.set_ylim(spans[-1][1], spans[0][0])
    ax.set_xlim(0, xmax * 1.12)
    ax.set_xlabel("Relative abundance (%)", labelpad=4)
    ax.grid(axis="x", linestyle="--", linewidth=0.5, color="gray", alpha=0.6)
    ax.grid(False, axis="y")
    ax.set_axisbelow(True)

    for top, bottom, species in spans:
        ax.plot([-0.115, -0.115], [top + 0.1, bottom - 0.1], transform=ax.get_yaxis_transform(),
                color="black", linewidth=0.6, clip_on=False)
        ax.text(-0.14, (top + bottom) / 2, species, transform=ax.get_yaxis_transform(),
                rotation=90, ha="center", va="center", fontsize=9, style="italic")

    for spine in ax.spines.values():
        spine.set_linewidth(0.6)
        spine.set_color("black")

    handles = [
        Patch(facecolor=colors["oleate"], edgecolor="black", linewidth=0.4, label="Oleate"),
        Patch(facecolor=colors["stearate"], edgecolor="black", linewidth=0.4, label="Stearate"),
        Patch(facecolor="white", edgecolor="black", linewidth=0.4, hatch="////", label="Doubtful"),
    ]
    ax.legend(
        handles=handles, title=None, fontsize=8, frameon=False, loc="upper center",
        bbox_to_anchor=(0.5, 1.06), ncol=3, handlelength=1.0, columnspacing=1.0,
    )

    plt.tight_layout(pad=0.3)
    plt.savefig(path, format="svg", bbox_inches="tight")
    plt.savefig(Path(path).with_suffix(".png"), format="png", bbox_inches="tight", dpi=220)
    plt.close(fig)
    print("Chart saved to", path)


def main():
    table = build_table()
    print(f"{len(table)} highlighted proteins:", ", ".join(table["label"]))
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    plot_highlighted(table, OUTPUT_FILE)


if __name__ == '__main__':
    main()